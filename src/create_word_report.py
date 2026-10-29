"""
Script: src/create_word_report.py
Description: Generates a professionally formatted Microsoft Word (.docx) document
             of the entire 24-page AAT report matching BMSCE Department of Machine Learning
             guidelines, complete with embedded diagrams, plots, tables, and cover page.
"""

import os
import json
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE_DIR = os.path.dirname(os.path.dirname(__file__))
PLOTS_DIR = os.path.join(BASE_DIR, "plots")
DATA_DIR = os.path.join(BASE_DIR, "data")
RESULTS_DIR = os.path.join(BASE_DIR, "results")
OUTPUT_DOCX = os.path.join(BASE_DIR, "AAT_Report_Hallucination_Detection.docx")


def set_cell_background(cell, fill_hex):
    """Sets background shading of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell padding."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def add_styled_heading(doc, text, level):
    p = doc.add_heading(text, level=level)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    run = p.runs[0]
    if level == 1:
        run.font.name = "Calibri"
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    elif level == 2:
        run.font.name = "Calibri"
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
    elif level == 3:
        run.font.name = "Calibri"
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
    return p


def add_body_paragraph(doc, text, bold_prefix=None, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(space_after)
    if bold_prefix:
        r_b = p.add_run(bold_prefix)
        r_b.font.name = "Calibri"
        r_b.font.size = Pt(11)
        r_b.font.bold = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    return p


def add_bullet_point(doc, bold_prefix, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(4)
    r_b = p.add_run(bold_prefix)
    r_b.font.name = "Calibri"
    r_b.font.size = Pt(11)
    r_b.font.bold = True
    r = p.add_run(text)
    r.font.name = "Calibri"
    r.font.size = Pt(11)
    return p


def add_image_with_caption(doc, image_filename, caption_text, width_inches=5.8):
    img_path = os.path.join(PLOTS_DIR, image_filename)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        run = p_img.add_run()
        run.add_picture(img_path, width=Inches(width_inches))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(9.5)
        r_cap.font.italic = True
        r_cap.font.bold = True
        r_cap.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)


def format_table(table, header_bg="1E3A8A", alt_bg="F3F4F6"):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # Format header row
    for cell in table.rows[0].cells:
        set_cell_background(cell, header_bg)
        set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Calibri"
                r.font.size = Pt(10)
                r.font.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Format body rows
    for r_idx, row in enumerate(table.rows[1:]):
        bg = alt_bg if r_idx % 2 == 1 else "FFFFFF"
        for cell in row.cells:
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.name = "Calibri"
                    r.font.size = Pt(9.5)


def build_full_report():
    print("[*] Creating professionally formatted Word Report (.docx)...")
    doc = Document()

    # Set 1 inch margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # ==================== COVER PAGE ====================
    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_inst = p_inst.add_run("B. M. S. COLLEGE OF ENGINEERING\n")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(18)
    r_inst.font.bold = True
    r_inst.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    r_sub = p_inst.add_run("Autonomous Institute, Affiliated to VTU, Belagavi\n\nDEPARTMENT OF MACHINE LEARNING\nProgram: B.E. In Artificial Intelligence and Machine Learning\nAcademic Year: 2026 - 27, Session: July 2026 - Dec 2026\n\n")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.bold = True

    p_course = doc.add_paragraph()
    p_course.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_c = p_course.add_run("Natural Language Processing for LLM\n(24AM7PCLLM)\n")
    r_c.font.name = "Calibri"
    r_c.font.size = Pt(15)
    r_c.font.bold = True
    r_c.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)

    r_aat = p_course.add_run("Alternative Assessment Tool (AAT) Report\n\n")
    r_aat.font.name = "Calibri"
    r_aat.font.size = Pt(13)
    r_aat.font.bold = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_t = p_title.add_run("Topic: Automated Hallucination & Factual Consistency Detection in LLM Generations using Lexical Overlap + Logistic Regression and DistilRoBERTa Cross-Encoder\n\n")
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(13)
    r_t.font.bold = True
    r_t.font.color.rgb = RGBColor(0x43, 0x38, 0xCA)

    p_subm = doc.add_paragraph()
    p_subm.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_s = p_subm.add_run("Submitted by\n")
    r_s.font.bold = True
    r_s.font.size = Pt(11)

    # Student details table
    tbl_students = doc.add_table(rows=3, cols=2)
    tbl_students.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_students.rows[0].cells[0].paragraphs[0].add_run("Student Name:\nTeam Member 1:\nTeam Member 2:")
    tbl_students.rows[0].cells[1].paragraphs[0].add_run("\n[Team Member 1 Name]\n[Team Member 2 Name]")
    tbl_students.rows[1].cells[0].paragraphs[0].add_run("USN:\nTeam Member 1:\nTeam Member 2:")
    tbl_students.rows[1].cells[1].paragraphs[0].add_run("\n1BM23AI...\n1BM23AI...")
    tbl_students.rows[2].cells[0].paragraphs[0].add_run("Semester & Section:")
    tbl_students.rows[2].cells[1].paragraphs[0].add_run("7 B")
    for r in tbl_students.rows:
        for c in r.cells:
            set_cell_margins(c, top=80, bottom=80, left=120, right=120)

    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(30)

    # Valuation Report table
    p_val = doc.add_paragraph()
    p_val.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_val = p_val.add_run("Valuation Report\n(to be filled by faculty)")
    r_val.font.bold = True
    r_val.font.italic = True
    r_val.font.color.rgb = RGBColor(0xDC, 0x26, 0x26)

    tbl_val = doc.add_table(rows=4, cols=2)
    tbl_val.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_val.rows[0].cells[0].paragraphs[0].add_run("Score:")
    tbl_val.rows[0].cells[1].paragraphs[0].add_run("")
    tbl_val.rows[1].cells[0].paragraphs[0].add_run("Comments:")
    tbl_val.rows[1].cells[1].paragraphs[0].add_run("")
    tbl_val.rows[2].cells[0].paragraphs[0].add_run("Faculty In-charge:")
    tbl_val.rows[2].cells[1].paragraphs[0].add_run("Dr Monika Puttaramaiah")
    tbl_val.rows[3].cells[0].paragraphs[0].add_run("Faculty Signature with date:")
    tbl_val.rows[3].cells[1].paragraphs[0].add_run("")
    for r in tbl_val.rows:
        for c in r.cells:
            set_cell_margins(c, top=100, bottom=100, left=120, right=120)

    doc.add_page_break()

    # ==================== TABLE OF CONTENTS ====================
    add_styled_heading(doc, "Table of Contents", level=1)
    toc_items = [
        ("1. Introduction", "1"),
        ("   1.1 Background of Natural Language Processing", "1"),
        ("   1.2 Importance of the Selected NLP Topic", "1"),
        ("   1.3 Objective of the Study", "1"),
        ("   1.4 Scope of the Report", "2"),
        ("2. Selected NLP Topic Overview", "3"),
        ("   2.1 Description of the NLP Task", "3"),
        ("   2.2 Real-World Use Cases", "3"),
        ("   2.3 Relevance to the Syllabus and Industry", "4"),
        ("   2.4 Key Challenges in This NLP Task", "4"),
        ("3. Real-Time Application / API Identified", "5"),
        ("   3.1 Application or API Name", "5"),
        ("   3.2 Purpose and Functionality", "5"),
        ("   3.3 Architecture or Platform on Which It Is Built", "5"),
        ("   3.4 Input/Output Behavior", "6"),
        ("   3.5 Limitations and Practical Considerations", "6"),
        ("4. NLP Model / Architecture Study", "7"),
        ("   4.1 Theoretical Foundations of the Model", "7"),
        ("   4.2 Algorithm(s) Used", "7"),
        ("   4.3 Architecture Diagram / Layer-Wise Design", "7"),
        ("   4.4 Data Flow in the Model", "8"),
        ("   4.5 Mathematical Formulation", "9"),
        ("   4.6 Why This Model Suits the Selected NLP Task", "9"),
        ("5. Dataset and Preprocessing", "10"),
        ("   5.1 Dataset Description (Source, Size, Structure)", "10"),
        ("   5.2 Data Cleaning Steps", "11"),
        ("   5.3 Tokenization, Lemmatization, Stopwords Removal", "11"),
        ("   5.4 Train/Test Split Strategy", "12"),
        ("   5.5 Feature Engineering or Embedding Approach", "12"),
        ("6. Implementation Methodology", "13"),
        ("   6.1 Tools, Libraries, and Environment Used", "13"),
        ("   6.2 Model Implementation Steps", "13"),
        ("   6.3 API Integration", "14"),
        ("   6.4 Pseudocode / Workflow Diagram", "14"),
        ("   6.5 Hyperparameters and Training Configuration", "15"),
        ("7. Results and Evaluation", "16"),
        ("   7.1 Performance Metrics", "16"),
        ("   7.2 Visualization of Results", "16"),
        ("   7.3 Qualitative and Quantitative Analysis", "17"),
        ("   7.4 Error Analysis", "18"),
        ("   7.5 Comparison Between Approaches", "19"),
        ("8. Discussion", "20"),
        ("   8.1 Interpretation of Results", "20"),
        ("   8.2 Strengths and Weaknesses of the Model", "20"),
        ("   8.3 Challenges Encountered During Implementation", "20"),
        ("   8.4 Impact of Dataset Quality and Preprocessing", "21"),
        ("   8.5 How API or Application Enhances the NLP Task", "21"),
        ("9. Conclusion", "22"),
        ("   9.1 Summary of the Work Done", "22"),
        ("   9.2 Key Learnings and Insights", "22"),
        ("   9.3 Real-World Applicability", "22"),
        ("   9.4 Future Improvements or Extensions", "22"),
        ("10. References", "24")
    ]
    for title, pg in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(title)
        r1.font.name = "Calibri"
        r1.font.size = Pt(10)
        # Dots
        dots_count = max(5, 75 - len(title))
        r2 = p.add_run(" " + "." * dots_count + " " + pg)
        r2.font.name = "Calibri"
        r2.font.size = Pt(10)
        r2.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    doc.add_page_break()

    # ==================== CHAPTER 1 ====================
    add_styled_heading(doc, "1. Introduction", level=1)
    add_styled_heading(doc, "1.1 Background of Natural Language Processing", level=2)
    add_body_paragraph(doc, "Natural Language Processing (NLP) is the branch of Artificial Intelligence that enables computational systems to read, interpret, reason over, and synthesize human language. Early NLP systems in the 1960s to 1980s were predominantly rule-based, relying on handcrafted context-free grammars and curated semantic lexicons. These were succeeded by statistical machine learning methods—such as Naive Bayes, Support Vector Machines (SVM), and Logistic Regression—which learned predictive decision boundaries over sparse Bag-of-Words (BoW) and Term Frequency–Inverse Document Frequency (TF-IDF) feature spaces.")
    add_body_paragraph(doc, "Over the past decade, the field experienced a monumental paradigm shift towards deep neural architectures: static continuous word vectors (Word2Vec, GloVe), recurrent networks (RNN, LSTM), and ultimately the self-attention-based Transformer architecture introduced by Vaswani et al. in 2017. Transformers replace recurrence with multi-head self-attention, permitting every token in a sequence to directly attend to all other tokens in parallel. This pre-training and downstream fine-tuning paradigm gave rise to modern Large Language Models (LLMs) such as GPT, Llama, and Gemini. However, as generative LLMs become ubiquitous across enterprise workflows, they exhibit a severe vulnerability: hallucination—the tendency to generate fluent, authoritative-sounding statements that contradict or extrapolate beyond factual reference sources.")

    add_styled_heading(doc, "1.2 Importance of the Selected NLP Topic", level=2)
    add_body_paragraph(doc, "Factual consistency verification—evaluating whether an LLM generation logically adheres to a verified ground-truth context—is among the most critical deployment challenges in applied AI. In Retrieval-Augmented Generation (RAG) and automated clinical summarization, unverified generations pose severe operational and legal risks. Automated factual consistency detection enables organizations to:")
    add_bullet_point(doc, "i. Automated Guardrails: ", "Intercept and filter hallucinated generations before delivering them to users.")
    add_bullet_point(doc, "ii. RAG Validation: ", "Score and filter retrieved context-generation pairs to prevent misinformation propagation.")
    add_bullet_point(doc, "iii. Regulatory Compliance: ", "Audit clinical and legal summaries to eliminate fabricated dosages or citations.")
    add_bullet_point(doc, "iv. Benchmark Evaluation: ", "Quantify and monitor generative model fidelity over time across production pipelines.")

    add_styled_heading(doc, "1.3 Objective of the Study", level=2)
    add_body_paragraph(doc, "The core objectives of this AAT project are:")
    add_bullet_point(doc, "i. ", "To formulate an end-to-end binary factual consistency and hallucination detection system for paired premise-hypothesis text.")
    add_bullet_point(doc, "ii. ", "To construct an interpretable classical machine learning baseline using Dense Lexical Overlap Features (Unigram/Bigram Jaccard, Token Coverage, Negation Discrepancy) with tuned Logistic Regression.")
    add_bullet_point(doc, "iii. ", "To evaluate a Transformer cross-encoder (DistilRoBERTa fine-tuned on NLI) on the exact same held-out test set under out-of-domain conditions.")
    add_bullet_point(doc, "iv. ", "To benchmark both models across accuracy, precision, recall, macro F1, parameter scale, training time, and inference latency.")
    add_bullet_point(doc, "v. ", "To analyze model behavior on adversarial linguistic phenomena including high-overlap entity swaps, numerical shifts, and polarity reversals.")
    add_bullet_point(doc, "vi. ", "To deploy both models in an interactive Streamlit web application for real-time inference and demonstration.")

    add_styled_heading(doc, "1.4 Scope of the Report", level=2)
    add_body_paragraph(doc, "This report covers the theoretical foundations, implementation methodology, empirical evaluation, and qualitative error analysis of both approaches. The experimental benchmark is scoped to binary classification (Class 1 = Factual / Grounded, Class 0 = Hallucinated / Unfaithful). To guarantee rigorous evaluation and avoid data leakage, an Out-of-Domain (OOD) stratified split is enforced across 25 isolated scientific, historical, and biomedical domains.")

    doc.add_page_break()

    # ==================== CHAPTER 2 ====================
    add_styled_heading(doc, "2. Selected NLP Topic Overview", level=1)
    add_styled_heading(doc, "2.1 Description of the NLP Task", level=2)
    add_body_paragraph(doc, "The selected NLP task is Natural Language Inference (NLI), formulated as a Sentence-Pair Factual Consistency Classification task. Given a reference passage C (premise) and an LLM-generated claim G (hypothesis), the model must predict label y in {0 = Hallucinated, 1 = Factual} along with a posterior confidence score P(y | C, G). It is a supervised classification problem where models evaluate directional semantic entailment between paired texts.")

    add_image_with_caption(doc, "fig1_workflow.png", "Figure 1: Overall workflow of the proposed factual consistency detection system")

    add_styled_heading(doc, "2.2 Real-World Use Cases", level=2)
    add_bullet_point(doc, "i. Enterprise RAG Gateways: ", "Audits generative responses against enterprise knowledge bases prior to user display.")
    add_bullet_point(doc, "ii. Clinical Record Summarization: ", "Validates physician clinical notes and EHR summaries to prevent medical malpractice.")
    add_bullet_point(doc, "iii. Legal Contract Analysis: ", "Audits automated clause syntheses against statutory language to prevent liability.")
    add_bullet_point(doc, "iv. Financial News Summarization: ", "Verifies fiscal metrics, revenues, and quarterly margins against regulatory 10-K filings.")
    add_bullet_point(doc, "v. Search Engine Answer Snippets: ", "Guarantees algorithmic snippets truthfully reflect indexed page content.")

    add_styled_heading(doc, "2.3 Relevance to the Syllabus and Industry", level=2)
    add_body_paragraph(doc, "The project maps directly to the curriculum of Natural Language Processing for LLM (24AM7PCLLM). It synthesizes vector space models, dense lexical overlap metrics, L2-regularized probabilistic linear classifiers, the Transformer architecture, multi-head self-attention mechanisms, and knowledge distillation. In industrial practice, lightweight lexical models and deep cross-encoders are deployed collaboratively in hierarchical cascades: fast lexical filters reject gross hallucinations in microseconds, reserving computationally intensive cross-encoders for high-overlap edge cases.")

    add_styled_heading(doc, "2.4 Key Challenges in This NLP Task", level=2)
    add_bullet_point(doc, "i. The Lexical Overlap Mirage: ", "Adversarial hallucinations reuse 80-95% of context words, easily deceiving bag-of-words classifiers.")
    add_bullet_point(doc, "ii. Negation and Antonyms: ", "Subtle verb polarity flips ('decreases' to 'increases') invert factual truth while maintaining high token overlap.")
    add_bullet_point(doc, "iii. Numerical & Date Precision: ", "Small alterations in dates or numbers constitute critical hallucinations that require character/subword granularity.")
    add_bullet_point(doc, "iv. Out-of-Domain Generalization: ", "The system must evaluate ungrounded claims across unseen scientific fields without domain retraining.")

    doc.add_page_break()

    # ==================== CHAPTER 3 ====================
    add_styled_heading(doc, "3. Real-Time Application / API Identified", level=1)
    add_styled_heading(doc, "3.1 Application or API Name", level=2)
    add_body_paragraph(doc, "The project uses the Hugging Face Transformers API and Model Hub (cross-encoder/nli-distilroberta-base) for deep cross-attention, delivered to end users through a custom Streamlit web application titled Real-Time LLM Hallucination & Factuality Detection.")

    add_styled_heading(doc, "3.2 Purpose and Functionality", level=2)
    add_body_paragraph(doc, "The application enables any researcher or user to paste a reference passage and an LLM generation to receive real-time verification:")
    add_bullet_point(doc, "i. Live Dual-Model Prediction: ", "Displays side-by-side classifications for Lexical ML and DistilRoBERTa.")
    add_bullet_point(doc, "ii. Consensus Badge: ", "Color-coded status indicating model agreement or flagging edge-case disagreements.")
    add_bullet_point(doc, "iii. Diagnostic Overlap Metrics: ", "Shows Jaccard similarity, token coverage %, bigram precision, and negation mismatch flags.")
    add_bullet_point(doc, "iv. Benchmarks & Evaluation Tab: ", "Renders the empirical metric comparison table, confusion matrices, and latency charts.")

    add_styled_heading(doc, "3.3 Architecture or Platform on Which It Is Built", level=2)
    add_body_paragraph(doc, "The system is written in Python 3.13 and organized into modular files:")

    tbl_proj = doc.add_table(rows=8, cols=2)
    proj_rows = [
        ("Module", "Responsibility"),
        ("src/dataset.py", "Generates 25-domain benchmark dataset, partitions OOD splits, generates EDA plots"),
        ("src/preprocessing.py", "Text normalization, contraction expansion, 8-dimensional dense lexical feature extraction"),
        ("src/baseline_model.py", "Hyperparameter tuning of C, training, and evaluation of Lexical Logistic Regression"),
        ("src/transformer_model.py", "Loads DistilRoBERTa Cross-Encoder checkpoint, runs CUDA batch evaluation on test split"),
        ("src/evaluation.py", "Builds comparison tables, compiles error_analysis.csv, and generates Figures 7, 8, 9, 10, 11"),
        ("src/predict.py", "HallucinationPredictor class: provides unified inference for both architectures"),
        ("app/app.py", "Interactive Streamlit web application featuring live detection and benchmark dashboards")
    ]
    for r_i, (c1, c2) in enumerate(proj_rows):
        tbl_proj.rows[r_i].cells[0].paragraphs[0].add_run(c1)
        tbl_proj.rows[r_i].cells[1].paragraphs[0].add_run(c2)
    format_table(tbl_proj)

    add_styled_heading(doc, "3.4 Input/Output Behavior", level=2)
    tbl_io = doc.add_table(rows=5, cols=3)
    io_rows = [
        ("Aspect", "Lexical Overlap + Logistic Regression", "DistilRoBERTa Cross-Encoder"),
        ("Input", "Paired text strings: Context and Claim", "Concatenated [CLS] Context [SEP] Claim [SEP]"),
        ("Internal Representation", "8-dimensional normalized dense feature vector", "Sequence of 768-d contextual embeddings"),
        ("Output", "Label (Fact/Hallu), P(Hallu), P(Fact)", "Label (Fact/Hallu), Softmax over 3 NLI logits"),
        ("Execution Platform", "Commodity CPU (Single thread)", "NVIDIA GeForce RTX 3050 Laptop GPU / CPU")
    ]
    for r_i, (c1, c2, c3) in enumerate(io_rows):
        tbl_io.rows[r_i].cells[0].paragraphs[0].add_run(c1)
        tbl_io.rows[r_i].cells[1].paragraphs[0].add_run(c2)
        tbl_io.rows[r_i].cells[2].paragraphs[0].add_run(c3)
    format_table(tbl_io)

    add_styled_heading(doc, "3.5 Limitations and Practical Considerations", level=2)
    add_bullet_point(doc, "i. Latency Discrepancy: ", "DistilRoBERTa takes 3.37 ms on GPU (~45 ms on CPU), which is ~2,592x slower than the baseline (0.0013 ms).")
    add_bullet_point(doc, "ii. Token Truncation: ", "Contexts exceeding 256 subwords are truncated, potentially omitting tail-end facts.")
    add_bullet_point(doc, "iii. Local Prototype: ", "The application is built for local auditing and lacks distributed queueing or auth layers.")

    doc.add_page_break()

    # ==================== CHAPTER 4 ====================
    add_styled_heading(doc, "4. NLP Model / Architecture Study", level=1)
    add_styled_heading(doc, "4.1 Theoretical Foundations of the Model", level=2)
    add_body_paragraph(doc, "The classical baseline leverages Lexical Semantics and Set-Theoretic Information Retrieval, quantifying factual consistency through continuous overlap ratios (Jaccard similarity, unigram/bigram token precision, length ratios) and syntactic negation mismatch indicators. In contrast, the Transformer model utilizes a Cross-Encoder Natural Language Inference (NLI) Architecture. Unlike bi-encoders which encode texts independently, cross-encoders perform full bidirectional self-attention across both texts simultaneously, allowing tokens in the generation to attend directly to all tokens in the context.")

    add_styled_heading(doc, "4.2 Algorithm(s) Used", level=2)
    tbl_algo = doc.add_table(rows=6, cols=3)
    algo_rows = [
        ("Property", "Model 1: Lexical Baseline", "Model 2: Transformer Cross-Encoder"),
        ("Algorithm", "Dense Overlap + Logistic Regression", "DistilRoBERTa Cross-Encoder (NLI)"),
        ("Model Type", "Linear Discriminative Classifier", "Deep Bidirectional Transformer Encoder"),
        ("Features", "8-d normalized lexical overlap vector", "Word-level BPE contextual embeddings"),
        ("Parameters", "9 learnable weights (8 weights + bias)", "≈ 82 million parameters"),
        ("Optimization", "L-BFGS, max_iter = 1000, C = 0.01", "AdamW during pre-training/fine-tuning")
    ]
    for r_i, (c1, c2, c3) in enumerate(algo_rows):
        tbl_algo.rows[r_i].cells[0].paragraphs[0].add_run(c1)
        tbl_algo.rows[r_i].cells[1].paragraphs[0].add_run(c2)
        tbl_algo.rows[r_i].cells[2].paragraphs[0].add_run(c3)
    format_table(tbl_algo)

    add_styled_heading(doc, "4.3 Architecture Diagram / Layer-Wise Design", level=2)
    add_image_with_caption(doc, "fig2_baseline_pipeline.png", "Figure 2: Lexical Overlap + Logistic Regression baseline pipeline")
    add_image_with_caption(doc, "fig3_transformer_architecture.png", "Figure 3: Layer-wise architecture of DistilRoBERTa Cross-Encoder")

    add_styled_heading(doc, "4.4 Data Flow in the Model", level=2)
    add_body_paragraph(doc, "Model 1: (C, G) -> clean_text() -> extract 8-d feature vector -> StandardScaler() -> w·x + b -> sigmoid -> probability -> threshold at 0.5.")
    add_body_paragraph(doc, "Model 2: (C, G) -> BPE Tokenizer -> [CLS] Context [SEP] Claim [SEP] -> 6 Transformer Encoder Blocks -> [CLS] pooling -> Linear head -> Softmax -> argmax(P(Entailment) vs P(Contradiction)).")

    add_styled_heading(doc, "4.5 Mathematical Formulation", level=2)
    add_body_paragraph(doc, "Lexical Unigram and Bigram Jaccard Similarity:")
    add_body_paragraph(doc, "J(C, G) = |T_C ∩ T_G| / |T_C ∪ T_G|,     J_bi(C, G) = |B_C ∩ B_G| / |B_C ∪ B_G|")
    add_body_paragraph(doc, "Logistic Regression probability and L2-regularized objective:")
    add_body_paragraph(doc, "P(y = 1 | x) = σ(w^T x + b) = 1 / (1 + e^-(w^T x + b))")
    add_body_paragraph(doc, "L(w) = (1/2)||w||^2 + C · Σ_i [ -y_i log p_i - (1 - y_i) log(1 - p_i) ]")
    add_body_paragraph(doc, "Scaled Dot-Product Attention in DistilRoBERTa layers:")
    add_body_paragraph(doc, "Attention(Q, K, V) = softmax( (Q K^T) / sqrt(d_k) ) V")
    add_body_paragraph(doc, "MultiHead(Q, K, V) = Concat(head_1, ..., head_12) W^O")

    add_styled_heading(doc, "4.6 Why This Model Suits the Selected NLP Task", level=2)
    add_bullet_point(doc, "i. Resolves Entity Swaps: ", "Self-attention links mismatched subjects and objects across sentence boundaries.")
    add_bullet_point(doc, "ii. Robust to High Overlap: ", "Identifies contradictory assertions even when 90% of words are identical.")
    add_bullet_point(doc, "iii. Distillation Efficiency: ", "DistilRoBERTa offers 97% of full RoBERTa accuracy at 2x the inference speed.")

    doc.add_page_break()

    # ==================== CHAPTER 5 ====================
    add_styled_heading(doc, "5. Dataset and Preprocessing", level=1)
    add_styled_heading(doc, "5.1 Dataset Description (Source, Size, Structure)", level=2)
    add_body_paragraph(doc, "The dataset comprises 5,000 paired records constructed across 25 diverse scientific, historical, medical, and financial domains. Each instance consists of context, generation, binary label, and domain category.")

    tbl_ds = doc.add_table(rows=7, cols=2)
    ds_rows = [
        ("Property", "Value"),
        ("Source", "Curated Multi-Domain Scientific & Historical Factuality Benchmark"),
        ("Total Text Pairs", "5,000 balanced pairs"),
        ("Classes", "2 – Hallucinated (0), Factual (1)"),
        ("Mean Generation Length", "20.3 words (137 characters)"),
        ("Median Generation Length", "20 words (range: 13 to 26 words)"),
        ("Mean Context Length", "26.3 words (182 characters)")
    ]
    for r_i, (c1, c2) in enumerate(ds_rows):
        tbl_ds.rows[r_i].cells[0].paragraphs[0].add_run(c1)
        tbl_ds.rows[r_i].cells[1].paragraphs[0].add_run(c2)
    format_table(tbl_ds)

    add_image_with_caption(doc, "fig4_class_distribution.png", "Figure 4: Class distribution of the sampled dataset")
    add_image_with_caption(doc, "fig5_length_distribution.png", "Figure 5: Generation word-count distribution by grounding class")

    add_styled_heading(doc, "5.2 Data Cleaning Steps", level=2)
    add_body_paragraph(doc, "The clean_text() function strips HTML tags, expands contractions, standardizes whitespace, and preserves critical punctuation marks (. ! ?).")

    add_styled_heading(doc, "5.3 Tokenization, Lemmatization, Stopwords Removal", level=2)
    add_bullet_point(doc, "i. Stopwords Removal: ", "Deliberately NOT applied. Removing stopwords strips negation words ('not', 'never') and alters truth polarity.")
    add_bullet_point(doc, "ii. Lemmatization: ", "Not applied, as verb tense alterations represent factual hallucinations.")
    add_bullet_point(doc, "iii. Tokenization: ", "DistilRoBERTa uses Byte-Pair Encoding (BPE, 50,265 vocabulary) preserving subwords.")

    add_styled_heading(doc, "5.4 Train/Test Split Strategy", level=2)
    tbl_split = doc.add_table(rows=5, cols=5)
    split_rows = [
        ("Split", "Pairs", "Hallucinated", "Factual", "Domain Purpose"),
        ("Train", "4,000", "2,000", "2,000", "20 Base Knowledge Domains"),
        ("Validation", "500", "250", "250", "3 Unseen Domains (Tuning C)"),
        ("Test (Held-out)", "500", "250", "250", "2 Held-out Domains (LIGO, Base Editing)"),
        ("Total", "5,000", "2,500", "2,500", "Strict Out-of-Domain Partitioning")
    ]
    for r_i, row in enumerate(split_rows):
        for c_i, val in enumerate(row):
            tbl_split.rows[r_i].cells[c_i].paragraphs[0].add_run(val)
    format_table(tbl_split)

    add_styled_heading(doc, "5.5 Feature Engineering or Embedding Approach", level=2)
    add_body_paragraph(doc, "Model 1 uses an 8-dimensional dense feature vector capturing unigram/bigram Jaccard, precision, recall, negation mismatch, and length ratios. Model 2 uses learned 768-d contextual embeddings.")

    doc.add_page_break()

    # ==================== CHAPTER 6 ====================
    add_styled_heading(doc, "6. Implementation Methodology", level=1)
    add_styled_heading(doc, "6.1 Tools, Libraries, and Environment Used", level=2)
    tbl_tools = doc.add_table(rows=7, cols=3)
    tools_rows = [
        ("Category", "Tool / Library", "Use in Project"),
        ("Language", "Python 3.13.7", "Entire pipeline implementation"),
        ("Data Handling", "pandas 2.3.3, NumPy 2.4.4", "Dataset management and array operations"),
        ("Classical ML", "scikit-learn 1.8.0", "StandardScaler, LogisticRegression, classification_report"),
        ("Deep Learning", "PyTorch 2.6.0+cu124, Transformers 5.21.0", "DistilRoBERTa inference, CUDA GPU execution"),
        ("Visualization", "Matplotlib 3.10.9, Seaborn 0.13.2", "Confusion matrices, metric charts, EDA plots"),
        ("Deployment", "Streamlit 1.61.1", "Real-time dual-model web application")
    ]
    for r_i, row in enumerate(tools_rows):
        for c_i, val in enumerate(row):
            tbl_tools.rows[r_i].cells[c_i].paragraphs[0].add_run(val)
    format_table(tbl_tools)

    add_styled_heading(doc, "6.2 Model Implementation Steps", level=2)
    add_bullet_point(doc, "1. ", "Execute src/dataset.py to generate 5,000 paired instances and produce Figures 4 and 5.")
    add_bullet_point(doc, "2. ", "Extract 8-d dense feature matrix in src/baseline_model.py; tune C on validation; train final model.")
    add_bullet_point(doc, "3. ", "Run src/transformer_model.py to evaluate DistilRoBERTa on 500 test pairs under CUDA acceleration.")
    add_bullet_point(doc, "4. ", "Execute src/evaluation.py to compile error_analysis.csv and generate Figures 7, 8, 9, 10, 11.")
    add_bullet_point(doc, "5. ", "Run tests/test_cases.py to probe models on hand-crafted rubric test cases.")
    add_bullet_point(doc, "6. ", "Launch Streamlit application with 'streamlit run app/app.py'.")

    add_styled_heading(doc, "6.3 API Integration", level=2)
    add_body_paragraph(doc, "Transformers AutoTokenizer and AutoModelForSequenceClassification APIs load 'cross-encoder/nli-distilroberta-base'. The model is wrapped in HallucinationPredictor with st.cache_resource for low-latency web serving.")

    add_styled_heading(doc, "6.4 Hyperparameters and Training Configuration", level=2)
    tbl_hyp = doc.add_table(rows=7, cols=2)
    hyp_rows = [
        ("Hyperparameter / Setting", "Configured Value"),
        ("Baseline Feature Dimension", "8 dense continuous lexical features"),
        ("Baseline Solver / C", "L-BFGS, max_iter = 1000, C = 0.01 (tuned on validation F1 = 0.7067)"),
        ("Transformer Checkpoint", "cross-encoder/nli-distilroberta-base (82M params)"),
        ("Max Sequence Length", "256 tokens (truncated/padded)"),
        ("Batch Size & Device", "16 pairs per batch, NVIDIA GeForce RTX 3050 Laptop GPU"),
        ("Decision Threshold", "Argmax: Factual if P(Entailment) >= P(Contradiction)")
    ]
    for r_i, row in enumerate(hyp_rows):
        for c_i, val in enumerate(row):
            tbl_hyp.rows[r_i].cells[c_i].paragraphs[0].add_run(val)
    format_table(tbl_hyp)

    doc.add_page_break()

    # ==================== CHAPTER 7 ====================
    add_styled_heading(doc, "7. Results and Evaluation", level=1)
    add_styled_heading(doc, "7.1 Performance Metrics", level=2)
    add_body_paragraph(doc, "Both models were evaluated on the exact same 500 held-out test pairs across unseen knowledge domains:")

    tbl_res = doc.add_table(rows=10, cols=4)
    res_rows = [
        ("Metric", "Lexical + Logistic Regression", "DistilRoBERTa Cross-Encoder", "Difference"),
        ("Accuracy", "82.80%", "94.00%", "+11.20%"),
        ("Macro Precision", "83.16%", "94.64%", "+11.48%"),
        ("Macro Recall", "82.80%", "94.00%", "+11.20%"),
        ("Macro F1-score", "82.75%", "93.98%", "+11.23%"),
        ("Weighted F1-score", "82.75%", "93.98%", "+11.23%"),
        ("Hallucinated Class (0) F1", "81.86%", "94.34%", "+12.48%"),
        ("Factual Class (1) F1", "83.65%", "93.62%", "+9.97%"),
        ("Training Time", "0.01 seconds", "Pre-trained checkpoint", "—"),
        ("Inference Latency", "0.0013 ms / pair", "3.37 ms / pair", "≈ 2,592x slower")
    ]
    for r_i, row in enumerate(res_rows):
        for c_i, val in enumerate(row):
            tbl_res.rows[r_i].cells[c_i].paragraphs[0].add_run(val)
    format_table(tbl_res)

    add_styled_heading(doc, "7.2 Visualization of Results", level=2)
    add_image_with_caption(doc, "fig7_confusion_matrices.png", "Figure 7: Confusion matrices: Lexical + Logistic Regression (left) and DistilRoBERTa (right)")

    tbl_cm = doc.add_table(rows=3, cols=5)
    cm_rows = [
        ("Model", "TN (True Hallu)", "FP (False Fact)", "FN (Missed Fact)", "TP (True Fact)"),
        ("Lexical + Logistic Regression", "194", "56", "30", "220"),
        ("DistilRoBERTa Cross-Encoder", "250", "0", "30", "220")
    ]
    for r_i, row in enumerate(cm_rows):
        for c_i, val in enumerate(row):
            tbl_cm.rows[r_i].cells[c_i].paragraphs[0].add_run(val)
    format_table(tbl_cm)

    add_image_with_caption(doc, "fig8_overall_metrics.png", "Figure 8: Comparison of overall metrics")
    add_image_with_caption(doc, "fig9_classwise_metrics.png", "Figure 9: Class-wise precision, recall, and F1-score")

    add_styled_heading(doc, "7.3 Qualitative and Quantitative Analysis", level=2)
    add_body_paragraph(doc, "Quantitative: DistilRoBERTa improves test accuracy by +11.20% (94.00% vs 82.80%). The primary driver of this gain is the elimination of False Positives: the baseline made 56 false-positive errors on high-overlap hallucinations, whereas DistilRoBERTa achieved 0 false positives (100% precision on the factual class).")
    add_body_paragraph(doc, "Qualitative: A specialized rubric test suite of four challenging cases evaluated model behavior on subtle linguistic shifts:")

    tbl_rubric = doc.add_table(rows=5, cols=5)
    rubric_rows = [
        ("Case", "Category", "Expected", "Baseline (LR)", "DistilRoBERTa"),
        ("TC-1", "Clearly factual paraphrase", "Factual", "Factual (69.0%)", "Factual (52.9%)"),
        ("TC-2", "Direct contradiction (Polarity flip)", "Hallucinated", "Factual (74.5%) [Error]", "Hallucinated (99.9%)"),
        ("TC-3", "High-overlap entity swap", "Hallucinated", "Hallucinated (95.6%)", "Hallucinated (99.3%)"),
        ("TC-4", "Temporal / Date fabrication", "Hallucinated", "Hallucinated (87.1%)", "Hallucinated (99.8%)")
    ]
    for r_i, row in enumerate(rubric_rows):
        for c_i, val in enumerate(row):
            tbl_rubric.rows[r_i].cells[c_i].paragraphs[0].add_run(val)
    format_table(tbl_rubric)

    add_body_paragraph(doc, "In TC-2 ('Metformin increases glucose production'), the generation reuses 63.6% of context tokens. The classical baseline is deceived by the high token overlap and predicts Factual (74.5% confidence). DistilRoBERTa's self-attention captures the semantic inversion of the verb, predicting Hallucinated (99.9% confidence).")

    add_styled_heading(doc, "7.4 Error Analysis", level=2)
    add_image_with_caption(doc, "fig10_model_agreement.png", "Figure 10: Prediction agreement between the two models on the test set")
    add_bullet_point(doc, "i. Both correct: ", "414 pairs (82.8%). Both models agree on standard paraphrases and low-overlap hallucinations.")
    add_bullet_point(doc, "ii. Only Transformer correct: ", "56 pairs (11.2%). Corresponds to adversarial entity swaps where baseline is misled by high overlap.")
    add_bullet_point(doc, "iii. Both wrong: ", "30 pairs (6.0%). Shared false negatives caused by highly abstract paraphrasing falling into NLI 'neutral'.")

    add_styled_heading(doc, "7.5 Comparison Between Approaches", level=2)
    add_image_with_caption(doc, "fig11_latency_comparison.png", "Figure 11: Inference latency per pair (ms, log scale)")

    tbl_comp = doc.add_table(rows=7, cols=3)
    comp_rows = [
        ("Criterion", "Lexical + Logistic Regression", "DistilRoBERTa Cross-Encoder"),
        ("Accuracy", "82.80%", "94.00% (Superior)"),
        ("Negation & Entity Handling", "Fails on subtle inversions", "Strong via cross-attention"),
        ("Training Cost", "0.01 seconds on CPU", "Pre-trained checkpoint (Hours on GPU)"),
        ("Inference Speed", "0.0013 ms / pair (Faster)", "3.37 ms / pair on GPU"),
        ("Model Size", "< 10 KB (.joblib file)", "≈ 320 MB (82M parameters)"),
        ("Hardware Needed", "Any commodity CPU", "CUDA GPU recommended for batch workloads")
    ]
    for r_i, row in enumerate(comp_rows):
        for c_i, val in enumerate(row):
            tbl_comp.rows[r_i].cells[c_i].paragraphs[0].add_run(val)
    format_table(tbl_comp)

    doc.add_page_break()

    # ==================== CHAPTER 8 ====================
    add_styled_heading(doc, "8. Discussion", level=1)
    add_styled_heading(doc, "8.1 Interpretation of Results", level=2)
    add_body_paragraph(doc, "The results clearly demonstrate that while lexical surface similarity provides a strong, high-throughput baseline (82.80% accuracy), it exhibits an inherent architectural blind spot on adversarial hallucinations. The 11.20 percentage-point improvement achieved by DistilRoBERTa confirms that factual consistency is governed by contextual relational binding rather than bag-of-words presence.")

    add_styled_heading(doc, "8.2 Strengths and Weaknesses of the Model", level=2)
    tbl_sw = doc.add_table(rows=3, cols=3)
    sw_rows = [
        ("Model", "Strengths", "Weaknesses"),
        ("Lexical + LR", "Instantaneous training (0.01 s); microsecond latency; ultra-lightweight; highly interpretable.", "Fails on high-overlap fabrications; ignores syntactic relationships; blind to semantic nuances."),
        ("DistilRoBERTa", "Superior accuracy (94.00%); robust against subtle entity/date swaps; zero false positives.", "2,592x slower inference; requires GPU; 256-token context truncation limit; black-box opacity.")
    ]
    for r_i, row in enumerate(sw_rows):
        for c_i, val in enumerate(row):
            tbl_sw.rows[r_i].cells[c_i].paragraphs[0].add_run(val)
    format_table(tbl_sw)

    add_styled_heading(doc, "8.3 Challenges Encountered During Implementation", level=2)
    add_bullet_point(doc, "i. Out-of-Domain Leakage: ", "Addressed topic memorization by isolating 25 knowledge domains strictly across splits.")
    add_bullet_point(doc, "ii. NLI Decision Mapping: ", "Resolved 3-class NLI outputs by formulating binary decision rule P(Entailment) >= P(Contradiction).")
    add_bullet_point(doc, "iii. GPU Tensor Batching: ", "Optimized PyTorch batching on NVIDIA RTX 3050 GPU, cutting latency down to 3.37 ms.")

    add_styled_heading(doc, "8.4 Impact of Dataset Quality and Preprocessing", level=2)
    add_body_paragraph(doc, "Preserving negation tokens while expanding contractions was critical for baseline performance. Omitting stopword removal prevented grammatical polarity inversion.")

    add_styled_heading(doc, "8.5 How API or Application Enhances the NLP Task", level=2)
    add_body_paragraph(doc, "The Streamlit application democratizes hallucination verification, allowing non-technical auditors to paste text, view consensus badges, and diagnose lexical overlap.")

    doc.add_page_break()

    # ==================== CHAPTER 9 ====================
    add_styled_heading(doc, "9. Conclusion", level=1)
    add_styled_heading(doc, "9.1 Summary of the Work Done", level=2)
    add_body_paragraph(doc, "This project built and compared two factual consistency classifiers for LLM generations. A balanced benchmark of 5,000 paired texts across 25 domains was partitioned under out-of-domain conditions. The classical baseline reached 82.80% accuracy in 0.0013 ms, while DistilRoBERTa Cross-Encoder achieved 94.00% accuracy and 93.98% macro F1, eliminating all false positives. Both models were deployed in an interactive Streamlit application.")

    add_styled_heading(doc, "9.2 Key Learnings and Insights", level=2)
    add_bullet_point(doc, "1. ", "Lexical overlap is insufficient for factuality verification due to adversarial token reuse.")
    add_bullet_point(doc, "2. ", "Cross-encoder self-attention successfully detects entity swaps, dates, and polarity inversions.")
    add_bullet_point(doc, "3. ", "A two-stage cascade combining fast lexical filtering with secondary transformer verification provides an optimal production trade-off.")

    add_styled_heading(doc, "9.3 Real-World Applicability", level=2)
    add_body_paragraph(doc, "The developed system is directly applicable as a verification firewall for RAG pipelines, clinical record summarization systems, and automated legal contract review.")

    add_styled_heading(doc, "9.4 Future Improvements or Extensions", level=2)
    add_bullet_point(doc, "i. Long-Context Chunking: ", "Implement sliding-window attention for multi-page documents.")
    add_bullet_point(doc, "ii. Model Quantization: ", "Export DistilRoBERTa to INT8 ONNX format for accelerated CPU inference.")
    add_bullet_point(doc, "iii. Token-Level Attribution: ", "Incorporate Integrated Gradients to highlight exact hallucinated spans.")

    doc.add_page_break()

    # ==================== CHAPTER 10 ====================
    add_styled_heading(doc, "10. References", level=1)
    refs = [
        "[1] A. Vaswani et al., 'Attention Is All You Need,' in Advances in Neural Information Processing Systems (NeurIPS), 2017.",
        "[2] Y. Liu et al., 'RoBERTa: A Robustly Optimized BERT Pretraining Approach,' arXiv:1907.11692, 2019.",
        "[3] V. Sanh et al., 'DistilBERT, a distilled version of BERT: smaller, faster, cheaper and lighter,' NeurIPS Workshop, 2019.",
        "[4] W. Kryscinski et al., 'Evaluating the Factual Consistency of Abstractive Text Summarization,' in Proc. EMNLP, 2020, pp. 9332–9346.",
        "[5] T. Goyal and G. Durrett, 'Evaluating Factuality in Generation with Dependency-level Entailment,' in Proc. EMNLP, 2020.",
        "[6] A. Wang et al., 'GLUE: A Multi-Task Benchmark and Analysis Platform for Natural Language Understanding,' in Proc. EMNLP, 2018.",
        "[7] N. Reimers and I. Gurevych, 'Sentence-BERT: Sentence Embeddings using Siamese BERT-Networks,' in Proc. EMNLP, 2019.",
        "[8] F. Pedregosa et al., 'Scikit-learn: Machine Learning in Python,' Journal of Machine Learning Research, vol. 12, pp. 2825–2830, 2011.",
        "[9] T. Wolf et al., 'Transformers: State-of-the-Art Natural Language Processing,' in Proc. EMNLP System Demos, 2020.",
        "[10] D. Jurafsky and J. H. Martin, Speech and Language Processing, 3rd ed. draft, Stanford University, 2024.",
        "[11] H. Zhang et al., 'HaluEval: A Large-Scale Hallucination Evaluation Benchmark for Large Language Models,' in Proc. EMNLP, 2023.",
        "[12] Streamlit Inc., 'Streamlit Documentation,' https://docs.streamlit.io, 2026."
    ]
    for r_text in refs:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(r_text)
        run.font.name = "Calibri"
        run.font.size = Pt(10)

    doc.save(OUTPUT_DOCX)
    print(f"[+] Complete Word Report generated successfully at: {OUTPUT_DOCX}")


if __name__ == "__main__":
    build_full_report()

